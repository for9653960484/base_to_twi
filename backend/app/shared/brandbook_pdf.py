"""Формирование PDF технологической карты (корпоративный стиль брендбука)."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from fpdf import FPDF
from fpdf.fonts import FontFace

from app.shared.brandbook_export import MAINTENANCE_LABELS_RU

BRAND_RGB = (26, 86, 142)
MUTED_RGB = (102, 102, 102)


def _cyrillic_font_path() -> Path:
    candidates = [
        Path(__file__).resolve().parent.parent / "assets" / "fonts" / "DejaVuSans.ttf",
        Path(r"C:\Windows\Fonts\arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
    ]
    for path in candidates:
        if path.is_file():
            return path
    raise FileNotFoundError(
        "Не найден шрифт с поддержкой кириллицы для PDF. "
        "Установите DejaVu Sans или используйте Windows."
    )


def _control_text(control: dict[str, Any]) -> str:
    if not control:
        return "—"
    return "; ".join(f"{k}: {v}" for k, v in control.items())


def render_tech_card_pdf(
    *,
    equipment_name: str,
    title: str,
    maintenance_type: str,
    work_items: list[dict[str, Any]],
    template_path: Path | None = None,
) -> bytes:
    """PDF как у листа безопасности: корпоративный DOCX-шаблон или письмо на бланке."""
    generated_at = datetime.now(timezone.utc).strftime("%d.%m.%Y")
    blocks = {
        "equipment_name": equipment_name or "—",
        "title": title or "Технологическая карта",
        "maintenance_type": MAINTENANCE_LABELS_RU.get(maintenance_type, maintenance_type),
        "generated_at": generated_at,
        "work_items": work_items or [],
        "work_items_text": _work_items_text(work_items or []),
    }
    if template_path and template_path.is_file() and template_path.suffix.lower() == ".docx":
        return _render_from_template(template_path, blocks)
    return _render_letter(blocks)


def _work_items_text(work_items: list[dict[str, Any]]) -> str:
    if not work_items:
        return "Не указаны"
    lines: list[str] = []
    for index, item in enumerate(work_items, start=1):
        order = item.get("order") or index
        description = str(item.get("description") or "—").strip()
        tools = ", ".join(item.get("tools") or []) or "—"
        safety = "; ".join(item.get("safety") or []) or "—"
        control = _control_text(item.get("control_params") or {})
        lines.append(
            f"{order}. {description}\n"
            f"Инструменты: {tools}\n"
            f"Безопасность: {safety}\n"
            f"Контроль: {control}"
        )
    return "\n".join(lines)


def _render_letter(blocks: dict[str, Any]) -> bytes:
    from app.shared.safety_pdf import _draw_footer, _draw_letterhead, _fonts, _section_title

    font_path, bold_path = _fonts()
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(15, 28, 15)
    pdf.add_font("Doc", "", str(font_path))
    if bold_path:
        pdf.add_font("Doc", "B", str(bold_path))
    pdf.add_page()
    _draw_letterhead(pdf, blocks["generated_at"])

    pdf.set_text_color(*BRAND_RGB)
    pdf.set_font("Doc", "B" if bold_path else "", 16)
    pdf.multi_cell(0, 8, "Технологическая карта", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    pdf.set_fill_color(240, 244, 248)
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Doc", "B" if bold_path else "", 12)
    pdf.multi_cell(
        0,
        8,
        f"Оборудование: {blocks['equipment_name']}",
        fill=True,
        new_x="LMARGIN",
        new_y="NEXT",
    )
    pdf.ln(1)
    pdf.set_font("Doc", "", 11)
    pdf.multi_cell(0, 6, blocks["title"], new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Doc", "", 11)
    pdf.multi_cell(
        0,
        6,
        f"Вид ТО: {blocks['maintenance_type']}",
        new_x="LMARGIN",
        new_y="NEXT",
    )
    pdf.ln(3)

    _section_title(pdf, "Перечень работ", bold_path)
    _work_table(pdf, blocks["work_items"])
    pdf.ln(8)
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Doc", "", 11)
    pdf.multi_cell(0, 7, "Составил: ________________    Дата: __________", new_x="LMARGIN", new_y="NEXT")
    pdf.multi_cell(0, 7, "Ознакомлен: ________________", new_x="LMARGIN", new_y="NEXT")

    _draw_footer(pdf)
    return bytes(pdf.output())


def _render_from_template(template_path: Path, blocks: dict[str, Any]) -> bytes:
    import io

    from docx import Document

    from app.shared.safety_pdf import (
        _document_text,
        _fonts,
        _header_image,
        _replace_placeholders,
        _section_title,
        _write_docx_table,
        _write_paragraph,
    )

    doc = Document(str(template_path))
    original = _document_text(doc)
    _replace_placeholders(
        doc,
        {
            "{{EQUIPMENT_NAME}}": blocks["equipment_name"],
            "{{TITLE}}": blocks["title"],
            "{{MAINTENANCE_TYPE}}": blocks["maintenance_type"],
            "{{GENERATED_AT}}": blocks["generated_at"],
            "{{WORK_ITEMS_BLOCK}}": blocks["work_items_text"],
            "{{WORK_ITEMS_TABLE}}": blocks["work_items_text"],
        },
    )

    font_path, bold_path = _fonts()
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(15, 15, 15)
    pdf.add_font("Doc", "", str(font_path))
    if bold_path:
        pdf.add_font("Doc", "B", str(bold_path))
    pdf.add_page()

    image = _header_image(doc)
    if image:
        pdf.image(io.BytesIO(image), x=15, y=12, w=40)
        pdf.set_y(28)

    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            _write_paragraph(pdf, paragraph.text, bold=True, bold_path=bold_path, size=11, color=BRAND_RGB)

    if "{{EQUIPMENT_NAME}}" not in original:
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("Doc", "B" if bold_path else "", 12)
        pdf.multi_cell(
            0,
            7,
            f"Оборудование: {blocks['equipment_name']}",
            new_x="LMARGIN",
            new_y="NEXT",
        )
        pdf.ln(1)

    for paragraph in doc.paragraphs:
        heading = (paragraph.style.name or "").startswith("Heading") if paragraph.style else False
        _write_paragraph(
            pdf,
            paragraph.text,
            bold=heading,
            bold_path=bold_path,
            size=14 if heading else 11,
            color=BRAND_RGB if heading else (0, 0, 0),
        )

    for table in doc.tables:
        _write_docx_table(pdf, table)

    if "{{WORK_ITEMS_BLOCK}}" not in original and "{{WORK_ITEMS_TABLE}}" not in original:
        pdf.ln(2)
        _section_title(pdf, "Перечень работ", bold_path)
        _work_table(pdf, blocks["work_items"])
    if "{{MAINTENANCE_TYPE}}" not in original:
        pdf.ln(1)
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("Doc", "", 11)
        pdf.multi_cell(
            0,
            6,
            f"Вид ТО: {blocks['maintenance_type']}",
            new_x="LMARGIN",
            new_y="NEXT",
        )

    return bytes(pdf.output())


def _work_table(pdf: FPDF, work_items: list[dict[str, Any]]) -> None:
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Doc", "", 9)
    if not work_items:
        pdf.multi_cell(0, 6, "Не указаны", new_x="LMARGIN", new_y="NEXT")
        return
    with pdf.table(
        width=180,
        col_widths=(10, 62, 38, 38, 32),
        line_height=5,
        text_align="LEFT",
        headings_style=FontFace(fill_color=(240, 244, 248)),
    ) as table:
        header = table.row()
        for text in ("№", "Описание работ", "Инструменты", "Безопасность", "Контроль"):
            header.cell(text)
        for index, item in enumerate(work_items, start=1):
            row = table.row()
            row.cell(str(item.get("order") or index))
            row.cell(str(item.get("description") or "—"))
            row.cell(", ".join(item.get("tools") or []) or "—")
            row.cell("; ".join(item.get("safety") or []) or "—")
            row.cell(_control_text(item.get("control_params") or {}))
