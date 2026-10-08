"""PDF листа техники безопасности: корпоративный шаблон или бланк письма."""

from __future__ import annotations

import io
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from fpdf import FPDF
from fpdf.fonts import FontFace

from app.core.config import settings
from app.shared.brandbook_export import BRAND_PRIMARY
from app.shared.brandbook_pdf import BRAND_RGB, MUTED_RGB, _cyrillic_font_path

def build_sample_safety_template() -> bytes:
    """Образец DOCX-шаблона письма с полями подстановки."""
    doc = Document()
    title = doc.add_heading(settings.APP_NAME, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = BRAND_PRIMARY

    subtitle = doc.add_paragraph("Корпоративный бланк")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("{{GENERATED_AT}}").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    heading = doc.add_heading("Лист по технике безопасности", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    equipment = doc.add_paragraph()
    equipment.add_run("Оборудование: ").bold = True
    equipment.add_run("{{EQUIPMENT_NAME}}")

    name = doc.add_paragraph()
    name.add_run("Наименование: ").bold = True
    name.add_run("{{TITLE}}")

    doc.add_heading("Средства индивидуальной защиты", level=2)
    doc.add_paragraph("{{PPE_BLOCK}}")
    doc.add_heading("Условия и правила работы", level=2)
    doc.add_paragraph("{{CONDITIONS_BLOCK}}")
    doc.add_heading("Обязательные правила", level=2)
    doc.add_paragraph("{{NOTES}}")

    footer = doc.add_paragraph("Документ сформирован системой Base To")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def render_safety_pdf(
    *,
    equipment_name: str,
    title: str,
    ppe: list[dict[str, Any]],
    work_conditions: dict[str, Any],
    notes: str | None,
    template_path: Path | None,
) -> bytes:
    generated_at = datetime.now(timezone.utc).strftime("%d.%m.%Y")
    blocks = _blocks(equipment_name, title, ppe, work_conditions, notes, generated_at)
    if template_path and template_path.is_file() and template_path.suffix.lower() == ".docx":
        return _render_from_template(template_path, blocks)
    return _render_letter(blocks)


def _blocks(
    equipment_name: str,
    title: str,
    ppe: list[dict[str, Any]],
    work_conditions: dict[str, Any],
    notes: str | None,
    generated_at: str,
) -> dict[str, Any]:
    return {
        "equipment_name": equipment_name or "—",
        "title": title or "Лист по технике безопасности",
        "generated_at": generated_at,
        "ppe": ppe or [],
        "ppe_text": _ppe_text(ppe),
        "conditions": _condition_rows(work_conditions),
        "conditions_text": _conditions_text(work_conditions),
        "notes": (notes or "").strip() or "—",
    }


def _ppe_text(ppe: list[dict[str, Any]]) -> str:
    if not ppe:
        return "Не указаны"
    lines = []
    for index, item in enumerate(ppe, start=1):
        name = str(item.get("name") or "").strip()
        purpose = str(item.get("purpose") or "").strip()
        mandatory = "обязательно" if item.get("mandatory", True) else "необязательно"
        line = f"{index}. {name}"
        if purpose:
            line += f" — {purpose}"
        line += f" ({mandatory})"
        lines.append(line)
    return "\n".join(lines)


def _condition_rows(work_conditions: dict[str, Any]) -> list[tuple[str, str]]:
    data = work_conditions or {}
    rows: list[tuple[str, str]] = []
    for label, key in (
        ("Температура", "temperature"),
        ("Влажность", "humidity"),
        ("Напряжение", "voltage"),
    ):
        value = str(data.get(key) or "").strip()
        if value:
            rows.append((label, value))
    for item in data.get("other") or []:
        if not isinstance(item, dict):
            continue
        name = str(item.get("name") or "").strip()
        value = str(item.get("value") or "").strip()
        if name and value:
            rows.append((name, value))
    return rows


def _conditions_text(work_conditions: dict[str, Any]) -> str:
    rows = _condition_rows(work_conditions)
    if not rows:
        return "Не указаны"
    return "\n".join(f"{label}: {value}" for label, value in rows)


def _render_letter(blocks: dict[str, Any]) -> bytes:
    font_path, bold_path = _fonts()
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(15, 28, 15)
    pdf.add_font("Doc", "", str(font_path))
    if bold_path:
        pdf.add_font("Doc", "B", str(bold_path))
    pdf.add_page()
    _draw_letterhead(pdf, blocks["generated_at"])

    pdf.set_text_color(*BRAND_RGB)
    pdf.set_font("Doc", "B" if bold_path else "", 16)
    pdf.multi_cell(0, 8, "Лист по технике безопасности", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    pdf.set_fill_color(240, 244, 248)
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Doc", "B" if bold_path else "", 12)
    pdf.multi_cell(
        0,
        8,
        f"Оборудование: {blocks['equipment_name']}",
        fill=True,
        new_x="LMARGIN",
        new_y="NEXT",
    )
    pdf.ln(1)
    pdf.set_font("Doc", "", 11)
    pdf.multi_cell(0, 6, blocks["title"], new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    _section_title(pdf, "Средства индивидуальной защиты", bold_path)
    _ppe_table(pdf, blocks["ppe"])
    pdf.ln(3)
    _section_title(pdf, "Условия и правила работы", bold_path)
    _conditions_table(pdf, blocks["conditions"])
    pdf.ln(3)
    _section_title(pdf, "Обязательные правила", bold_path)
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Doc", "", 11)
    pdf.multi_cell(0, 6, blocks["notes"], new_x="LMARGIN", new_y="NEXT")
    pdf.ln(8)
    pdf.set_font("Doc", "", 11)
    pdf.multi_cell(0, 7, "Составил: ________________    Дата: __________", new_x="LMARGIN", new_y="NEXT")
    pdf.multi_cell(0, 7, "Ознакомлен: ________________", new_x="LMARGIN", new_y="NEXT")

    _draw_footer(pdf)
    return bytes(pdf.output())


def _render_from_template(template_path: Path, blocks: dict[str, Any]) -> bytes:
    doc = Document(str(template_path))
    original = _document_text(doc)
    mapping = {
        "{{EQUIPMENT_NAME}}": blocks["equipment_name"],
        "{{TITLE}}": blocks["title"],
        "{{GENERATED_AT}}": blocks["generated_at"],
        "{{PPE_BLOCK}}": blocks["ppe_text"],
        "{{CONDITIONS_BLOCK}}": blocks["conditions_text"],
        "{{NOTES}}": blocks["notes"],
    }
    _replace_placeholders(doc, mapping)

    font_path, bold_path = _fonts()
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(15, 15, 15)
    pdf.add_font("Doc", "", str(font_path))
    if bold_path:
        pdf.add_font("Doc", "B", str(bold_path))
    pdf.add_page()

    image = _header_image(doc)
    if image:
        stream = io.BytesIO(image)
        pdf.image(stream, x=15, y=12, w=40)
        pdf.set_y(28)

    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            _write_paragraph(pdf, paragraph.text, bold=True, bold_path=bold_path, size=11, color=BRAND_RGB)

    if "{{EQUIPMENT_NAME}}" not in original:
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("Doc", "B" if bold_path else "", 12)
        pdf.multi_cell(
            0,
            7,
            f"Оборудование: {blocks['equipment_name']}",
            new_x="LMARGIN",
            new_y="NEXT",
        )
        pdf.ln(1)

    for paragraph in doc.paragraphs:
        heading = (paragraph.style.name or "").startswith("Heading") if paragraph.style else False
        _write_paragraph(
            pdf,
            paragraph.text,
            bold=heading,
            bold_path=bold_path,
            size=14 if heading else 11,
            color=BRAND_RGB if heading else (0, 0, 0),
        )

    for table in doc.tables:
        _write_docx_table(pdf, table)

    if "{{PPE_BLOCK}}" not in original:
        pdf.ln(2)
        _section_title(pdf, "Средства индивидуальной защиты", bold_path)
        _ppe_table(pdf, blocks["ppe"])
    if "{{CONDITIONS_BLOCK}}" not in original:
        pdf.ln(2)
        _section_title(pdf, "Условия и правила работы", bold_path)
        _conditions_table(pdf, blocks["conditions"])
    if "{{NOTES}}" not in original and blocks["notes"] != "—":
        pdf.ln(2)
        _section_title(pdf, "Обязательные правила", bold_path)
        pdf.set_text_color(0, 0, 0)
        pdf.set_font("Doc", "", 11)
        pdf.multi_cell(0, 6, blocks["notes"], new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())


def _document_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(parts)


def _replace_placeholders(doc: Document, mapping: dict[str, str]) -> None:
    targets = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                targets.extend(cell.paragraphs)
    for section in doc.sections:
        targets.extend(section.header.paragraphs)
        targets.extend(section.footer.paragraphs)
    for paragraph in targets:
        text = paragraph.text
        updated = text
        for key, value in mapping.items():
            updated = updated.replace(key, value)
        if updated != text:
            paragraph.text = updated


def _header_image(doc: Document) -> bytes | None:
    for section in doc.sections:
        for rel in section.header.part.rels.values():
            if "image" in rel.reltype:
                return rel.target_part.blob
    return None


def _write_paragraph(
    pdf: FPDF,
    text: str,
    *,
    bold: bool,
    bold_path: Path | None,
    size: int,
    color: tuple[int, int, int],
) -> None:
    cleaned = (text or "").strip()
    if not cleaned:
        return
    pdf.set_text_color(*color)
    pdf.set_font("Doc", "B" if bold and bold_path else "", size)
    for line in cleaned.splitlines() or [cleaned]:
        pdf.multi_cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(1)


def _write_docx_table(pdf: FPDF, table) -> None:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return
    width = 180
    col_count = max(len(row) for row in rows)
    col_width = width / col_count
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Doc", "", 9)
    with pdf.table(
        width=width,
        col_widths=tuple(col_width for _ in range(col_count)),
        line_height=5,
        text_align="LEFT",
        headings_style=FontFace(fill_color=(240, 244, 248)),
    ) as grid:
        for row_values in rows:
            row = grid.row()
            for index in range(col_count):
                value = row_values[index] if index < len(row_values) else ""
                row.cell(value or "—")
    pdf.ln(2)


def _ppe_table(pdf: FPDF, ppe: list[dict[str, Any]]) -> None:
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Doc", "", 10)
    if not ppe:
        pdf.multi_cell(0, 6, "Не указаны", new_x="LMARGIN", new_y="NEXT")
        return
    with pdf.table(
        width=180,
        col_widths=(12, 58, 78, 32),
        line_height=5,
        text_align="LEFT",
        headings_style=FontFace(fill_color=(240, 244, 248)),
    ) as table:
        header = table.row()
        for text in ("№", "Наименование", "Назначение", "Обязательность"):
            header.cell(text)
        for index, item in enumerate(ppe, start=1):
            row = table.row()
            row.cell(str(index))
            row.cell(str(item.get("name") or "—"))
            row.cell(str(item.get("purpose") or "—"))
            row.cell("обязательно" if item.get("mandatory", True) else "необязательно")


def _conditions_table(pdf: FPDF, rows: list[tuple[str, str]]) -> None:
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Doc", "", 10)
    if not rows:
        pdf.multi_cell(0, 6, "Не указаны", new_x="LMARGIN", new_y="NEXT")
        return
    with pdf.table(
        width=180,
        col_widths=(60, 120),
        line_height=5,
        text_align="LEFT",
        headings_style=FontFace(fill_color=(240, 244, 248)),
    ) as table:
        header = table.row()
        header.cell("Параметр")
        header.cell("Значение")
        for label, value in rows:
            row = table.row()
            row.cell(label)
            row.cell(value)


def _section_title(pdf: FPDF, text: str, bold_path: Path | None) -> None:
    pdf.set_text_color(*BRAND_RGB)
    pdf.set_font("Doc", "B" if bold_path else "", 12)
    pdf.cell(0, 8, text, new_x="LMARGIN", new_y="NEXT")


def _draw_letterhead(pdf: FPDF, generated_at: str) -> None:
    pdf.set_fill_color(*BRAND_RGB)
    pdf.rect(0, 0, 210, 22, style="F")
    pdf.set_xy(15, 5)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Doc", "", 14)
    pdf.cell(120, 7, settings.APP_NAME, new_x="LMARGIN", new_y="NEXT")
    pdf.set_x(15)
    pdf.set_font("Doc", "", 9)
    pdf.cell(120, 6, "База знаний по техобслуживанию")
    pdf.set_xy(140, 7)
    pdf.set_font("Doc", "", 10)
    pdf.cell(55, 8, generated_at, align="R")
    pdf.set_y(28)
    pdf.set_text_color(0, 0, 0)


def _draw_footer(pdf: FPDF) -> None:
    pdf.set_y(-15)
    pdf.set_text_color(*MUTED_RGB)
    pdf.set_font("Doc", "", 8)
    pdf.cell(0, 8, "Документ сформирован системой Base To", align="C")


def _fonts() -> tuple[Path, Path | None]:
    regular = _cyrillic_font_path()
    bold_candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        Path("/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
    ]
    bold = next((path for path in bold_candidates if path.is_file()), None)
    return regular, bold
